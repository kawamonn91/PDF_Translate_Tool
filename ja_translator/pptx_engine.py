"""PowerPoint のテキストを読み取り、訳文に差し替える。図形・画像・表の配置はそのまま残す。

単位は「テキスト枠1つ」(表はセル1つ)。枠の中の段落は、訳文の改行ごとに段落として書き直し、
先頭の段落の書式(箇条書き・字間・色)を引き継ぐ。
"""

from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

from .fit import fit_text
from .fonts import measure
from .model import TextUnit
from .translate import needs_translation

EMU_PER_PT = 12700
DEFAULT_SIZE = 18.0
_ALIGN = {"CENTER": "center", "RIGHT": "right"}


def _frame_text(text_frame) -> str:
    return "\n".join("".join(run.text for run in p.runs).strip() for p in text_frame.paragraphs if p.runs).strip()


def _base_run(text_frame):
    for p in text_frame.paragraphs:
        for run in p.runs:
            if run.text.strip():
                return run
    return None


def _run_size(run) -> float:
    if run is not None and run.font.size is not None:
        return run.font.size.pt
    return DEFAULT_SIZE


def _run_color(run) -> tuple[float, float, float]:
    try:
        rgb = run.font.color.rgb if run is not None and run.font.color and run.font.color.type is not None else None
    except AttributeError:
        rgb = None
    if rgb is None:
        return (0.0, 0.0, 0.0)
    return (rgb[0] / 255, rgb[1] / 255, rgb[2] / 255)


def _inner_box(text_frame, width_pt: float, height_pt: float) -> tuple[float, float]:
    w = width_pt - (text_frame.margin_left.pt + text_frame.margin_right.pt)
    h = height_pt - (text_frame.margin_top.pt + text_frame.margin_bottom.pt)
    return max(w, 10.0), max(h, 10.0)


def _iter_frames(shapes, path: tuple[int, ...] = (), origin=(0.0, 0.0)):
    """(場所, 枠の位置と大きさpt, text_frame) を読み取り順に返す。グループ・表の中も辿る"""
    for idx, shape in enumerate(shapes):
        here = path + (idx,)
        left = (shape.left or 0) / EMU_PER_PT
        top = (shape.top or 0) / EMU_PER_PT
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_frames(shape.shapes, here, (left, top))
        elif getattr(shape, "has_table", False) and shape.has_table:
            table = shape.table
            y = top
            for r, row in enumerate(table.rows):
                x = left
                for c, cell in enumerate(row.cells):
                    width = table.columns[c].width / EMU_PER_PT
                    height = row.height / EMU_PER_PT
                    yield (here, (r, c)), (x, y, width, height), cell.text_frame
                    x += width
                y += row.height / EMU_PER_PT
        elif shape.has_text_frame:
            rect = (left, top, (shape.width or 0) / EMU_PER_PT, (shape.height or 0) / EMU_PER_PT)
            yield (here, None), rect, shape.text_frame


def _unit_id(slide_no: int, here: tuple[int, ...], cell: tuple[int, int] | None) -> str:
    base = f"s{slide_no}-" + ".".join(str(i) for i in here)
    return base if cell is None else f"{base}-r{cell[0]}c{cell[1]}"


def extract_units(path: str | Path) -> list[TextUnit]:
    prs = Presentation(str(path))
    units: list[TextUnit] = []
    for sno, slide in enumerate(prs.slides):
        for (here, cell), (x, y, width_pt, height_pt), frame in _iter_frames(slide.shapes):
            source = _frame_text(frame)
            if not needs_translation(source):
                continue
            run = _base_run(frame)
            align = "left"
            for p in frame.paragraphs:
                if p.runs and p.alignment is not None:
                    align = _ALIGN.get(str(p.alignment).split(".")[-1].split(" ")[0].upper(), "left")
                    break
            w, h = _inner_box(frame, width_pt, height_pt)
            units.append(
                TextUnit(
                    id=_unit_id(sno + 1, here, cell),
                    location={"slide": sno, "path": list(here), "cell": list(cell) if cell else None},
                    source=source,
                    bbox=(0.0, 0.0, w, h),
                    size=_run_size(run),
                    color=_run_color(run),
                    align=align,
                    bold=bool(run and run.font.bold),
                    extra={"rect": [x, y, width_pt, height_pt]},
                )
            )
    return units


def _find_frame(prs, location: dict):
    slide = prs.slides[location["slide"]]
    shapes = slide.shapes
    target = None
    path = list(location["path"])
    for depth, idx in enumerate(path):
        shape = list(shapes)[idx]
        if depth == len(path) - 1:
            target = shape
        else:
            shapes = shape.shapes
    if location.get("cell"):
        r, c = location["cell"]
        return target.table.cell(r, c).text_frame
    return target.text_frame


def _write_frame(text_frame, lines: list[str], size: float) -> None:
    body = text_frame._txBody
    paragraphs = body.findall(qn("a:p"))
    template_p = next((p for p in paragraphs if p.findall(qn("a:r"))), paragraphs[0])
    template_r = template_p.find(qn("a:r"))
    for p in paragraphs:
        body.remove(p)
    for line in lines:
        new_p = deepcopy(template_p)
        for child in list(new_p):
            if child.tag in (qn("a:r"), qn("a:br"), qn("a:fld")):
                new_p.remove(child)
        run = deepcopy(template_r) if template_r is not None else etree.SubElement(new_p, qn("a:r"))
        rpr = run.find(qn("a:rPr"))
        if rpr is None:
            rpr = etree.Element(qn("a:rPr"))
            run.insert(0, rpr)
        rpr.set("lang", "ja-JP")
        rpr.set("sz", str(int(round(size * 100))))
        run.find(qn("a:t")).text = line
        end = new_p.find(qn("a:endParaRPr"))
        if end is not None:
            end.addprevious(run)
        else:
            new_p.append(run)
        body.append(new_p)


def render(src_path: str | Path, units: list[TextUnit], out_path: str | Path) -> None:
    prs = Presentation(str(src_path))
    for unit in units:
        if not unit.translation:
            continue
        frame = _find_frame(prs, unit.location)
        size = unit.size
        if unit.scale != 1.0:
            size = unit.size * unit.scale
        else:
            w, h = unit.bbox[2], unit.bbox[3]
            size = fit_text(unit.translation, w, h, unit.size, lambda t, s: measure(t, s, unit.bold)).font_size
        lines = unit.translation.split("\n")
        _write_frame(frame, lines, size)
    prs.save(str(out_path))
