import pytest
from pptx import Presentation
from pptx.util import Inches, Pt

from ja_translator.project import Project, sidecar_path


@pytest.fixture
def sample_pptx(tmp_path):
    path = tmp_path / "deck.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Quarterly Sales Report"

    box = slide.shapes.add_textbox(Inches(0.5), Inches(1.8), Inches(3), Inches(0.6))
    box.text_frame.word_wrap = True
    run = box.text_frame.paragraphs[0].add_run()
    run.text = "Revenue grew by twelve percent compared with the previous year and the outlook stays strong."
    run.font.size = Pt(20)

    table = slide.shapes.add_table(2, 2, Inches(5), Inches(1.8), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "Region"
    table.cell(0, 1).text = "Growth rate by region"

    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(0.5), Inches(4), Inches(3), Inches(0.5)).text_frame.text = "Footer note"
    prs.save(path)
    return path


def test_extracts_titles_tables_and_group_text(sample_pptx):
    sources = [u.source for u in Project.open(sample_pptx).units]
    assert "Quarterly Sales Report" in sources
    assert "Region" in sources
    assert "Footer note" in sources


def test_export_writes_japanese_and_keeps_shape_count(sample_pptx, tmp_path):
    project = Project.open(sample_pptx)
    for unit in project.units:
        unit.translation = "訳文です" if unit.source != "Region" else "地域"
    out = tmp_path / "deck_ja.pptx"
    project.export(out)

    before = len(Presentation(sample_pptx).slides[0].shapes)
    after = Presentation(out)
    slide = after.slides[0]
    assert len(slide.shapes) == before
    texts = [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame]
    assert any("訳文です" in t for t in texts)
    table = next(shape.table for shape in slide.shapes if shape.has_table)
    assert table.cell(0, 0).text == "地域"


def test_long_translation_is_shrunk_to_fit_the_box(sample_pptx, tmp_path):
    project = Project.open(sample_pptx)
    unit = next(u for u in project.units if u.source.startswith("Revenue"))
    unit.translation = "前年比で売上が増加しました。"
    out = tmp_path / "shrink.pptx"
    project.export(out)

    slide = Presentation(out).slides[0]
    box = next(s for s in slide.shapes if s.has_text_frame and "前年比" in s.text_frame.text)
    sizes = [r.font.size.pt for p in box.text_frame.paragraphs for r in p.runs if r.font.size]
    assert sizes and max(sizes) < 20
    assert not project.warnings()


def test_japanese_run_is_marked_as_ja_jp(sample_pptx, tmp_path):
    project = Project.open(sample_pptx)
    project.units[0].translation = "売上報告書"
    out = tmp_path / "lang.pptx"
    project.export(out)
    title = Presentation(out).slides[0].shapes.title
    rpr = title.text_frame.paragraphs[0].runs[0]._r.find(
        "{http://schemas.openxmlformats.org/drawingml/2006/main}rPr"
    )
    assert rpr.get("lang") == "ja-JP"


def test_sidecar_is_written_next_to_source(sample_pptx):
    project = Project.open(sample_pptx)
    project.units[0].translation = "売上報告書"
    project.save_sidecar()
    assert sidecar_path(sample_pptx).exists()
    assert Project.open(sample_pptx).units[0].translation == "売上報告書"
    sidecar_path(sample_pptx).unlink()
